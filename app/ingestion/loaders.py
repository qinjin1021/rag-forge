"""多格式文档加载器(工厂模式)。

工程化要点:
- 按扩展名自动分发到对应解析器, 新增格式只需注册一个解析函数
- 解析失败抛统一异常, 由上层管道捕获并记录, 单个文件失败不阻塞批量入库
"""
from pathlib import Path

from app.core.logging import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".xlsx"}


class DocumentLoadError(Exception):
    """文档加载失败"""


def _load_plain(path: Path) -> str:
    """纯文本/Markdown: 优先utf-8, 兼容gbk(中文Windows常见)"""
    for encoding in ("utf-8", "gbk", "utf-16"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentLoadError(f"无法识别文件编码: {path.name}")


def _load_docx(path: Path) -> str:
    """Word文档: 按段落读取, 表格按 行/单元格 展开为文本"""
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _load_pdf(path: Path) -> str:
    """PDF: 逐页提取文本(扫描件/图片型PDF不适用, 需接OCR, 此处为工程边界)"""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _load_xlsx(path: Path) -> str:
    """Excel: 每行拼接为一条记录, 保留列结构便于检索"""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)


# 解析器注册表: 扩展名 -> 解析函数。新增格式在此登记即可
_LOADERS = {
    ".txt": _load_plain,
    ".md": _load_plain,
    ".docx": _load_docx,
    ".pdf": _load_pdf,
    ".xlsx": _load_xlsx,
}


def load_document(path: Path) -> str:
    """加载单个文档为纯文本。不支持/解析失败时抛DocumentLoadError"""
    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise DocumentLoadError(f"不支持的文件格式: {suffix}(支持: {SUPPORTED_EXTENSIONS})")
    if not path.is_file():
        raise DocumentLoadError(f"文件不存在: {path}")
    try:
        text = loader(path)
    except DocumentLoadError:
        raise
    except Exception as exc:  # 统一收口第三方解析库的各类异常
        raise DocumentLoadError(f"解析失败 {path.name}: {exc}") from exc

    if not text.strip():
        raise DocumentLoadError(f"文档内容为空: {path.name}")
    logger.info("加载文档成功: %s (%d字符)", path.name, len(text))
    return text
